"""File source helpers: format sniff, revision fingerprint, xlsx sheet listing."""

from __future__ import annotations

import os
from pathlib import Path

from ..models import DATA_FORMATS

_EXT_TO_FORMAT = {"csv": "csv", "parquet": "parquet", "pq": "parquet", "xlsx": "xlsx"}


def detect_format(path: str | Path) -> str | None:
    ext = Path(path).suffix.lower().lstrip(".")
    return _EXT_TO_FORMAT.get(ext)


def file_revision(path: str | Path) -> str:
    """Content fingerprint for a single file: mtime(seconds)+size."""
    st = os.stat(path)
    return f"{int(st.st_mtime)}-{st.st_size}"


def list_sheets(path: str | Path) -> list[str]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def xlsx_row_count(path: str | Path, sheet: str) -> int:
    """Approximate data rows via openpyxl read_only dims (header row excluded)."""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True)
    try:
        ws = wb[sheet]
        return max((ws.max_row or 1) - 1, 0)
    finally:
        wb.close()


def formats_ok(path: str | Path) -> bool:
    return detect_format(path) in DATA_FORMATS
